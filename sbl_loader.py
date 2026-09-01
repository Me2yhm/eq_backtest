"""SBL source loading, normalization, and deterministic borrow selection."""

from __future__ import annotations

import hashlib
import json
import math
import os
from datetime import date, datetime
from itertools import chain
from pathlib import Path
from typing import Any, Iterator

import polars as pl
import pyarrow as pa
import pyarrow.parquet as pq
from loguru import logger

SBL_CACHE_VERSION = 1
_YADING_HEADERS = ("生效日期", "股票代码", "股票名称", "市场", "利率", "总数")
_YADING_MARKETS = frozenset({"SH.QFII", "SH.HK", "SZ.QFII", "SZ.HK"})
_SCHEMA = pa.schema([
    ("date", pa.date32()),
    ("symbol", pa.string()),
    ("provider", pa.string()),
    ("channel", pa.string()),
    ("available_shares", pa.float64()),
    ("annual_borrow_rate", pa.float64()),
    ("source_row", pa.int64()),
    ("provider_order", pa.int32()),
])


def _sources(raw_sources: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    if not raw_sources:
        return []
    if not isinstance(raw_sources, list):
        raise ValueError("short_borrow_sources must be a list")
    result: list[dict[str, Any]] = []
    for order, raw in enumerate(raw_sources):
        if not isinstance(raw, dict):
            raise ValueError(f"short_borrow_sources[{order}] must be a mapping")
        if set(raw).difference({"provider", "adapter", "path"}):
            raise ValueError("short_borrow_sources entries accept provider, adapter, and path only")
        provider, adapter, path = raw.get("provider"), raw.get("adapter"), raw.get("path")
        if not isinstance(provider, str) or not provider:
            raise ValueError(f"short_borrow_sources[{order}].provider must be a non-empty string")
        if adapter != "yading":
            raise ValueError("only the 'yading' short_borrow_sources adapter is currently supported")
        resolved = Path(path) if isinstance(path, (str, Path)) else None
        if resolved is None or resolved.suffix.lower() != ".xlsx" or not resolved.is_file():
            raise FileNotFoundError(f"Invalid Yading SBL path: {path!r}")
        result.append({"provider": provider, "adapter": adapter, "path": resolved, "provider_order": order})
    return result


def _signature(path: Path) -> dict[str, int | str]:
    stat = path.stat()
    return {"path": path.resolve().as_posix(), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}


def sbl_source_signature(raw_sources: list[dict[str, Any]] | None, selection: str) -> list[dict[str, Any]]:
    """Return SBL cache-identity inputs after validating configured sources."""
    if selection != "min_available_rate":
        raise ValueError("borrow_selection must be 'min_available_rate'")
    return [
        {**{key: source[key] for key in ("provider", "adapter", "provider_order")}, "file": _signature(source["path"])}
        for source in _sources(raw_sources)
    ]


def _value_date(value: Any, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as exc:
        raise ValueError(f"Invalid Yading effective date at {label}: {value!r}") from exc


def _value_symbol(value: Any, label: str) -> str:
    if isinstance(value, int):
        code = f"{value:06d}"
    elif isinstance(value, float) and value.is_integer():
        code = f"{int(value):06d}"
    else:
        code = str(value).strip().zfill(6)
    if len(code) != 6 or not code.isdigit():
        raise ValueError(f"Invalid Yading stock code at {label}: {value!r}")
    return code


def _nonnegative(value: Any, field: str, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid Yading {field} at {label}: {value!r}") from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"Invalid Yading {field} at {label}: {value!r}")
    return number


def _iter_yading(source: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError("Yading SBL loading requires openpyxl; run uv sync") from exc

    workbook = load_workbook(source["path"], read_only=True, data_only=True)
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            values = worksheet.iter_rows(values_only=True)
            first = next(values, None)
            if first is None:
                continue
            headers = tuple("" if value is None else str(value).strip() for value in first[:6])
            if sheet_index == 0:
                if headers != _YADING_HEADERS:
                    raise ValueError(f"Unexpected Yading header in '{source['path']}': {headers}")
                rows, first_row = values, 2
            elif headers == _YADING_HEADERS:
                rows, first_row = values, 2
            else:
                rows, first_row = chain((first,), values), 1

            for row_number, row in enumerate(rows, start=first_row):
                if not any(value is not None for value in row):
                    continue
                if len(row) < 6:
                    raise ValueError(f"Yading row has fewer than six columns: {worksheet.title}!{row_number}")
                label = f"{source['path'].name}:{worksheet.title}!{row_number}"
                market = str(row[3]).strip()
                if market not in _YADING_MARKETS:
                    raise ValueError(
                        f"Unsupported Yading 市场 at {label}: {market!r}; expected {sorted(_YADING_MARKETS)}"
                    )
                exchange, channel = market.split(".", 1)
                yield {
                    "date": _value_date(row[0], label),
                    "symbol": f"{_value_symbol(row[1], label)}{'.XSHG' if exchange == 'SH' else '.XSHE'}",
                    "provider": source["provider"],
                    "channel": channel,
                    "available_shares": _nonnegative(row[5], "总数", label),
                    "annual_borrow_rate": _nonnegative(row[4], "利率", label),
                    "source_row": row_number,
                    "provider_order": source["provider_order"],
                }
    finally:
        workbook.close()


def _cache_path(cache_dir: Path, sources: list[dict[str, Any]]) -> Path:
    payload = {
        "version": SBL_CACHE_VERSION,
        "sources": [
            {**{key: source[key] for key in ("provider", "adapter", "provider_order")}, "file": _signature(source["path"])}
            for source in sources
        ],
    }
    digest = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()[:16]
    return cache_dir / f"sbl_normalized_{digest}.parquet"


def _write_cache(path: Path, sources: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.parquet")
    writer: pq.ParquetWriter | None = None
    rows: list[dict[str, Any]] = []
    try:
        for source in sources:
            for row in _iter_yading(source):
                rows.append(row)
                if len(rows) == 100_000:
                    table = pa.Table.from_pylist(rows, schema=_SCHEMA)
                    writer = writer or pq.ParquetWriter(temporary, _SCHEMA, compression="zstd")
                    writer.write_table(table)
                    rows.clear()
        table = pa.Table.from_pylist(rows, schema=_SCHEMA)
        writer = writer or pq.ParquetWriter(temporary, _SCHEMA, compression="zstd")
        writer.write_table(table)
    finally:
        if writer is not None:
            writer.close()
    os.replace(temporary, path)
    logger.info("SBL normalized cache: wrote – {}", path.name)


def _select_available_rows(
    rows: Iterator[dict[str, Any]],
    *,
    start: date,
    end: date | None,
) -> pl.DataFrame:
    """Select the deterministic available record for each date-symbol in one pass."""
    selected: dict[tuple[date, str], dict[str, Any]] = {}
    for row in rows:
        row_date = row["date"]
        if row_date < start or (end is not None and row_date > end) or row["available_shares"] <= 0:
            continue
        key = (row_date, row["symbol"])
        rank = (row["annual_borrow_rate"], row["provider_order"], row["channel"], row["source_row"])
        existing = selected.get(key)
        if existing is None or rank < (
            existing["annual_borrow_rate"],
            existing["provider_order"],
            existing["channel"],
            existing["source_row"],
        ):
            selected[key] = row
    return (
        pl.DataFrame(
            [
                {
                    "date": row["date"],
                    "symbol": row["symbol"],
                    "borrow_rate": row["annual_borrow_rate"],
                    "borrow_provider": row["provider"],
                    "borrow_channel": row["channel"],
                    "borrow_available": True,
                }
                for row in selected.values()
            ],
            schema={
                "date": pl.Date,
                "symbol": pl.String,
                "borrow_rate": pl.Float64,
                "borrow_provider": pl.String,
                "borrow_channel": pl.String,
                "borrow_available": pl.Boolean,
            },
        )
        .sort(["date", "symbol"])
    )


def load_borrow_availability(
    raw_sources: list[dict[str, Any]] | None,
    *,
    start: str,
    end: str | None,
    cache_dir: Path,
    selection: str = "min_available_rate",
    use_cache: bool = True,
) -> pl.DataFrame:
    """Load one selected available borrow record for each exact date-symbol."""
    if selection != "min_available_rate":
        raise ValueError("borrow_selection must be 'min_available_rate'")
    sources = _sources(raw_sources)
    if not sources:
        raise ValueError("short_only and long_short modes require short_borrow_sources")
    path = _cache_path(cache_dir, sources)
    start_date = date.fromisoformat(start)
    end_date = date.fromisoformat(end) if end else None
    if not use_cache:
        logger.info("SBL normalized cache: disabled; streaming source workbooks")
        return _select_available_rows(
            chain.from_iterable(_iter_yading(source) for source in sources),
            start=start_date,
            end=end_date,
        )
    if not path.exists():
        _write_cache(path, sources)
    else:
        logger.info("SBL normalized cache: hit – {}", path.name)

    raw = pl.scan_parquet(path).filter(pl.col("date") >= pl.lit(start_date))
    if end_date:
        raw = raw.filter(pl.col("date") <= pl.lit(end_date))
    result = (
        raw.filter(pl.col("available_shares") > 0)
        .sort(["date", "symbol", "annual_borrow_rate", "provider_order", "channel", "source_row"])
        .group_by(["date", "symbol"], maintain_order=True)
        .agg([
            pl.col("annual_borrow_rate").first().alias("borrow_rate"),
            pl.col("provider").first().alias("borrow_provider"),
            pl.col("channel").first().alias("borrow_channel"),
        ])
        .with_columns(pl.lit(True).alias("borrow_available"))
        .collect()
        .sort(["date", "symbol"])
    )
    logger.info("SBL availability: {} selected date-symbol rows", result.height)
    return result
